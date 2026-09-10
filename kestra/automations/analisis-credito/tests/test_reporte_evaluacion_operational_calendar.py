from datetime import date, datetime
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import Mock, patch

from openpyxl import Workbook
from reporte_evaluacion_report.analysis import build_month_report
from reporte_evaluacion_report.core import MonthDataset, NovedadEvent
from reporte_evaluacion_comisiones.operational_calendar import detect_operational_calendar, verify_day
from reporte_evaluacion_comisiones.operational_calendar_excel import build_operational_calendar_sheet


def event(oid=1, state="RevisionRiesgo", stamp=datetime(2026, 7, 10, 10), user="vendedor"):
    return NovedadEvent.from_api_row([oid, stamp.date().isoformat(), f"[{state}]",
                                    stamp.strftime("%d/%m/%y %H:%M:%S ") + user,
                                    10, 10, 10, "Policia", "Pagada"])


class OperationalCalendarTests(TestCase):
    def detect(self, events, *, confirmed=None, error=None):
        records = [{"id": e.event_id, "date": e.created_at.isoformat(), "state": e.parsed_state,
                    "user": e.usuario_evento, "application": e.solicitud_oid} for e in events if e.created_at]
        def check(client, day, limit):
            if error:
                raise error
            return {"records": [e for e in records if e["date"][:10] == day.isoformat()], "count_before": 0, "count_after": 0}
        with patch("reporte_evaluacion_comisiones.operational_calendar.verify_day", side_effect=check):
            result = detect_operational_calendar(Mock(), [SimpleNamespace(month_value="2026-07", month_events=events)],
                                                 {date(2026,7,9): "Feriado"}, today=date(2026,9,10),
                                                 users=("evaluador",), confirmed=confirmed)
        day = next(d for d in result["days"] if d["date"] == "2026-07-10")
        return result, day

    def test_intake_is_not_work_and_candidate_is_not_deducted(self):
        result, day = self.detect([event()], confirmed={})
        self.assertEqual(day["status"], "Posible no trabajado")
        self.assertEqual(day["team_events"], 0)
        self.assertEqual(day["events"], 1)
        self.assertFalse(day["extra_excluded"])
        self.assertEqual(result["extra_excluded_dates"], [])

    def test_confirmed_closure_deducts_day_not_cases(self):
        result, day = self.detect([event()])
        self.assertEqual(day["status"], "Cierre confirmado")
        self.assertTrue(day["extra_excluded"])
        self.assertEqual(result["extra_excluded_dates"], ["2026-07-10"])

    def test_any_team_activity_prevents_false_closure_even_without_decisions(self):
        _, day = self.detect([event(user="evaluador")], confirmed={})
        self.assertEqual(day["status"], "Con actividad")
        self.assertEqual(day["decision_events"], 0)
        self.assertEqual(day["team_events"], 1)

    def test_unknown_decision_actor_requires_roster_review(self):
        _, day = self.detect([event(state="Revisar", user="nuevo_analista")], confirmed={})
        self.assertEqual(day["status"], "Actividad por atribuir")
        self.assertFalse(day["extra_excluded"])

    def test_api_failure_never_means_zero_work(self):
        result, day = self.detect([], confirmed={}, error=ConnectionError("private endpoint"))
        self.assertEqual(day["status"], "Sin datos completos")
        self.assertFalse(day["extra_excluded"])
        self.assertEqual(result["verifications"]["2026-07-10"], {"error_type": "ConnectionError"})

    def test_conflicting_confirmed_closure_is_visible(self):
        _, day = self.detect([event(state="Revisar", user="evaluador")])
        self.assertEqual(day["status"], "Confirmación en conflicto")
        self.assertTrue(day["extra_excluded"])

    def test_count_and_timestamps_are_checked(self):
        c = Mock()
        row = [1, "2026-07-10", "[Revisar]", "10/07/26 09:00:00 evaluador", 10, "evaluador"]
        c.evaluate.return_value = 1
        c.evaluate_list.return_value = [row]
        self.assertEqual(len(verify_day(c, date(2026,7,10), 20000)["records"]), 1)
        c.evaluate.side_effect = [1, 2]
        with self.assertRaises(ValueError):
            verify_day(c, date(2026,7,10), 20000)
        c.evaluate.side_effect = None
        row[3] = "11/07/26 09:00:00 evaluador"
        with self.assertRaises(ValueError):
            verify_day(c, date(2026,7,10), 20000)
        c.evaluate_list.return_value = [row, row]
        c.evaluate.return_value = 2
        with self.assertRaises(ValueError):
            verify_day(c, date(2026,7,10), 2)

    def test_malformed_month_is_not_a_closed_day(self):
        e = event()
        e.created_at = None
        _, day = self.detect([e], confirmed={})
        self.assertEqual(day["status"], "Sin datos completos")
        self.assertFalse(day["extra_excluded"])

    def test_only_first_response_changes_and_no_cases_are_removed(self):
        events = [event(1, stamp=datetime(2026,7,9,12)),
                  event(2, "Confirmada", datetime(2026,7,13,8,20), "evaluador"),
                  event(3, "A Transferir", datetime(2026,7,13,9), "tesoreria"),
                  event(4, "Pagada", datetime(2026,7,13,9,10), "tesoreria")]
        dataset = MonthDataset("2026-07", "Julio", 1, events, [10], events)
        original = build_month_report(dataset, log=lambda _: None, excluded_dates={date(2026,7,9)})
        revised = build_month_report(dataset, log=lambda _: None, excluded_dates={date(2026,7,9)},
                                     first_response_extra_excluded_dates={date(2026,7,10)})
        self.assertEqual(original.first_response[0]["minutos"], 560)
        self.assertEqual(revised.first_response[0]["minutos"], 20)
        self.assertEqual(revised.transfer, original.transfer)
        self.assertEqual(revised.end_to_end, original.end_to_end)
        self.assertEqual(revised.legajos_sample, original.legajos_sample)
        self.assertEqual(len(revised.first_response), len(original.first_response))

    def test_trace_sheet_is_secondary_and_exposes_confirmations(self):
        audit, _ = self.detect([event()])
        wb = Workbook()
        ws = build_operational_calendar_sheet(wb, audit)
        self.assertEqual(wb.sheetnames[-1], "Jornadas evaluación")
        self.assertIn("ExcepcionesJornadasEvaluacion", ws.tables)
        self.assertIn("HistoriaJornadasEvaluacion", ws.tables)
        self.assertIsInstance(ws["A10"].value, date)
        self.assertEqual(ws["B11"].value, "Posible no trabajado")
        self.assertIsNone(ws.freeze_panes)
