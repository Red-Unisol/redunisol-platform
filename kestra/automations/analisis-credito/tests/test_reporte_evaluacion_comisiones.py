from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from reporte_evaluacion_report.analysis import build_month_report, business_seconds_between
from reporte_evaluacion_report.core import MonthDataset, NovedadEvent, derive_month_seed
from reporte_evaluacion_report.excel import _build_state_duration_rows
from reporte_evaluacion_report.storage import SQLiteDatasetStore
from reporte_evaluacion_comisiones.calendar import national_holidays
from reporte_evaluacion_comisiones.core import (
    CommissionApiClient, Loan, commission_rate, evaluate_commissions, fetch_loans, previous_months,
)
from reporte_evaluacion_comisiones.excel import build_commission_sheet
from reporte_evaluacion_comisiones.kestra_entrypoint import atomic_publish, generate_report


def dataset(month: str, *, count: int = 1, across_holiday: bool = False) -> MonthDataset:
    year, number = map(int, month.split("-"))
    start = datetime(year, number, 3, 9)
    while start.weekday() >= 5:
        start += timedelta(days=1)
    end = start + timedelta(minutes=10)
    if across_holiday:
        start, end = datetime(2026, 7, 8, 16), datetime(2026, 7, 10, 9)
    events = []
    for oid in range(1, count + 1):
        for index, (timestamp, state) in enumerate([
            (start, "RevisionRiesgo"), (end, "Confirmada"),
            (end + timedelta(minutes=5), "A Transferir"), (end + timedelta(minutes=15), "Pagada"),
        ]):
            events.append(NovedadEvent.from_api_row([
                oid * 10 + index, timestamp.date().isoformat(), f"[{state}]",
                timestamp.strftime("%d/%m/%y %H:%M:%S") + " analista", oid, oid, oid, "LINEA CBU", "Pagada",
            ]))
    return MonthDataset(month, month, derive_month_seed(202510, month), events, list(range(1, count + 1)), events)


def sample_loan(month: str = "2026-08") -> Loan:
    return Loan(100, 1000, month + "-03", 1, "Pagada", "Vendedor habilitado", Decimal("1000000"), Decimal("1200000"), "LINEA CBU")


class CommissionsTests(unittest.TestCase):
    def test_confirmed_rate_boundaries(self):
        for value, expected in [("9.5", ".005"), ("10", ".005"), ("10.00001", ".003"), ("11", ".003"), ("11.00001", ".001")]:
            with self.subTest(value=value):
                self.assertEqual(commission_rate(Decimal(value), Decimal(10)), Decimal(expected))
        self.assertIsNone(commission_rate(Decimal(1), Decimal(0)))
        with self.assertRaises(ValueError):
            commission_rate(Decimal("NaN"), Decimal(10))

    def test_three_previous_months_across_year_boundary(self):
        self.assertEqual(previous_months("2026-01"), ["2025-10", "2025-11", "2025-12"])
        self.assertEqual(previous_months("2026-02"), ["2025-11", "2025-12", "2026-01"])

    def test_calendar_mandatory_only_and_exceptional_transfer(self):
        calendar = national_holidays([2025, 2026])
        for day in [date(2025, 10, 10), date(2026, 7, 9), date(2026, 4, 2), date(2026, 4, 3), date(2026, 6, 15), date(2026, 11, 23)]:
            self.assertIn(day, calendar)
        for day in [date(2025, 10, 12), date(2025, 4, 17), date(2025, 8, 15), date(2026, 7, 10), date(2026, 3, 23), date(2026, 12, 7)]:
            self.assertNotIn(day, calendar)
        self.assertNotIn("Jueves Santo", calendar[date(2026, 4, 2)])
        with self.assertRaisesRegex(ValueError, "Calendario"):
            national_holidays([2027])

    def test_calendar_changes_all_v2_metrics_but_preserves_v1_and_sampling(self):
        data = dataset("2026-07", count=35, across_holiday=True)
        exclusions = frozenset(national_holidays([2026]))
        old = build_month_report(data, log=lambda _: None)
        new = build_month_report(data, log=lambda _: None, excluded_dates=exclusions)
        self.assertEqual(old.summary["first_response"]["promedio_minutos"], 660)
        self.assertEqual(new.summary["first_response"]["promedio_minutos"], 120)
        self.assertEqual(new.summary["end_to_end"]["promedio_minutos"], 135)
        self.assertEqual(old.summary["end_to_end"]["promedio_minutos"], 675)
        self.assertEqual(new.summary["transfer"]["promedio_minutos"], 10)
        self.assertEqual(new.legajos_sample, old.legajos_sample)
        self.assertEqual(len(new.legajos_sample), 30)
        state_rows = _build_state_duration_rows(new, excluded_dates=exclusions)
        old_state_rows = _build_state_duration_rows(old)
        self.assertEqual(old_state_rows[0]["duracion_laboral_min"] - state_rows[0]["duracion_laboral_min"], 540)
        a, b = datetime(2026, 7, 8, 16), datetime(2026, 7, 10, 9)
        self.assertEqual(business_seconds_between(a, b), 660 * 60)

    def test_reference_is_simple_monthly_mean_not_weighted_by_cases(self):
        reports = [build_month_report(dataset(m), log=lambda _: None) for m in ["2026-05", "2026-06", "2026-07", "2026-08"]]
        for report, minutes, cases in zip(reports, [5, 10, 15, 11], [100, 1, 1, 10]):
            report.summary["first_response"]["promedio_minutos"] = minutes
            report.summary["first_response"]["cases"] = cases
        results = evaluate_commissions(reports, {"2026-08": [sample_loan()]}, ["2026-08"])
        self.assertEqual(results[1]["reference"], Decimal(10))
        self.assertEqual(results[1]["rate"], Decimal(".003"))
        self.assertEqual(results[1]["amount"], Decimal("600.00"))
        reports[0].summary["first_response"]["promedio_minutos"] = None
        self.assertIsNone(evaluate_commissions(reports, {"2026-08": []}, ["2026-08"])[1]["amount"])

    def test_loan_projection_matches_video_and_detects_truncation_or_drift(self):
        client = Mock()
        client.evaluate.return_value = 1
        client.evaluate_list.return_value = [[1, 1001, "2026-08-03", 3, "Pagada", "Gloria Fernandez", 730000, 857750, "CBU"]]
        rows = fetch_loans(client, "2026-08")
        self.assertEqual(rows[0].amount, Decimal(730000))
        self.assertEqual(rows[0].capital, Decimal(857750))
        args = client.evaluate_list.call_args.args
        self.assertIn("[FechaEmision] < #2026-09-01#", args[0])
        self.assertIn("'Martin Rodriguez'", args[0])
        self.assertIn("[Solicitud.Estado.Descripcion] = 'Pagada'", args[0])
        client.evaluate.return_value = 2
        with self.assertRaisesRegex(ValueError, "cantidad"):
            fetch_loans(client, "2026-08")
        client.evaluate.side_effect = [1, 2]
        with self.assertRaisesRegex(ValueError, "cambio"):
            fetch_loans(client, "2026-08")

    def test_rejects_duplicate_loan_ids_and_invalid_amounts(self):
        client = Mock()
        row = [1, 1001, "2026-08-03", 3, "Pagada", "Gloria Fernandez", 100, 120, "CBU"]
        client.evaluate.return_value = 2
        client.evaluate_list.return_value = [row, row]
        with self.assertRaisesRegex(ValueError, "duplicado"):
            fetch_loans(client, "2026-08")
        client.evaluate.return_value = 1
        for amount in [-1, float("nan"), float("inf")]:
            row[6] = amount
            client.evaluate_list.return_value = [row]
            with self.assertRaises(ValueError):
                fetch_loans(client, "2026-08")

    def test_strict_events_client_never_uses_truncated_rows(self):
        with patch("reporte_evaluacion_report.core.EvaluateApiClient.evaluate_list", return_value=[[1]]):
            with self.assertRaisesRegex(ValueError, "limite"):
                CommissionApiClient("https://example.invalid").evaluate_list("True", "type", "ID", 1)

    def test_publish_preserves_v1_and_every_v2_run(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            old = root / "analisis-credito/reporte-evaluacion/ultimo.xlsx"
            old.parent.mkdir(parents=True)
            old.write_bytes(b"original")
            sources = [root / name for name in ["book.xlsx", "data.sqlite", "manifest.json"]]
            for source in sources:
                source.write_bytes(b"first")
            now = datetime(2026, 9, 1, 8, 15)
            latest, first = atomic_publish(*sources, root, now)
            sources[0].write_bytes(b"second")
            _, second = atomic_publish(*sources, root, now)
            self.assertNotEqual(first, second)
            self.assertEqual(first.read_bytes(), b"first")
            self.assertEqual(latest.read_bytes(), b"second")
            self.assertEqual(old.read_bytes(), b"original")
            self.assertEqual(len(list(latest.parent.joinpath("datos").glob("*.sqlite"))), 2)

    def test_review_counts_are_independent_per_month_and_require_full_sample(self):
        months = ["2026-07", "2026-08"]
        reports = {m: build_month_report(dataset(m, count=n), log=lambda _: None) for m, n in zip(months, [7, 0])}
        refs = {m: r for r, m in enumerate(["2026-04", "2026-05", "2026-06", *months], 5)}
        sheet = build_commission_sheet(Workbook(), months, reports, {m: [] for m in months}, refs, {m: (5, 4) for m in months})
        for row, month in [(50, "2026-07"), (114, "2026-08")]:
            self.assertIn('COUNTIFS(RevisionLegajos[Mes],"'+month+'"', sheet[f"A{row}"].value)
            self.assertIn('COUNTIF(RevisionLegajos[Mes],"'+month+'"', sheet[f"G{row}"].value)
            self.assertIn("'Reglas'!$B$13", sheet[f"G{row + 7}"].value)
        self.assertIn("G50=7", sheet["G57"].value)
        self.assertIn("G114=0", sheet["G121"].value)
        self.assertEqual(len(sheet.data_validations.dataValidation), 0)
        self.assertEqual([item.id for item in sheet.row_breaks.brk], [27, 67, 91])
        # Baseline months can have samples while the requested month has none.
        empty = build_commission_sheet(Workbook(), ["2026-08"], reports, {"2026-08": []}, refs, {"2026-08": (5, 4)})
        self.assertEqual(empty["A50"].value, "=0")
        self.assertEqual(empty["G50"].value, "=0")

    def test_generation_fetches_baseline_and_keeps_manual_commission_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.dict("os.environ", {
                "REPORTE_EVALUACION_BASE_URL": "https://example.invalid",
                "REPORTS_ROOT": tmp, "REPORT_INPUT_FROM_MONTH": "2026-08", "REPORT_INPUT_TO_MONTH": "2026-08",
                "TRIGGER_BODY_JSON": "{}",
            }), patch("reporte_evaluacion_comisiones.kestra_entrypoint.CommissionApiClient"), \
                patch("reporte_evaluacion_comisiones.kestra_entrypoint.fetch_month_dataset", side_effect=lambda client, month, **kw: dataset(month, count=31)) as fetch, \
                patch("reporte_evaluacion_comisiones.kestra_entrypoint.fetch_loans", return_value=[sample_loan()]):
                result = generate_report(datetime(2026, 9, 1, 8, 15))
            self.assertEqual([call.args[1] for call in fetch.call_args_list], ["2026-05", "2026-06", "2026-07", "2026-08"])
            workbook = load_workbook(result["latest_path"])
            self.assertEqual(workbook.sheetnames[0], "Comisiones")
            self.assertEqual(len(workbook["Muestreo legajos"]["A"]) - 4, 30)
            self.assertNotIn("Objetivos y comisiones", workbook.sheetnames)
            sheet = workbook["Comisiones"]
            self.assertEqual(sheet.max_column, 8)
            self.assertIn('"Correcto"', sheet["A50"].value)
            self.assertIn('"Incorrecto"', sheet["C50"].value)
            self.assertEqual(sheet["E50"].value, "=G50-A50-C50")
            self.assertIn("COUNT(E60,E58)=2", sheet["E61"].value)
            self.assertIn("COUNT(G10,G20,G30,G40)=4", sheet["E60"].value)
            self.assertIn("E50=0", sheet["G57"].value)
            self.assertEqual(workbook["Reglas"]["B11"].value, 28)
            self.assertEqual(workbook["Reglas"]["B12"].value, 26)
            self.assertEqual(workbook["Reglas"]["B13"].value, 30)
            self.assertIn("C10<=A10", sheet["G16"].value)
            self.assertIn("ALCANZADO", sheet["G12"].value)
            review = workbook["Muestreo legajos"]
            self.assertEqual(workbook.sheetnames[1], "Muestreo legajos")
            self.assertEqual([review.cell(row, 5).value for row in range(5, 35)], ["A revisar"] * 30)
            self.assertTrue(all(review.cell(row, 2).value for row in range(5, 35)))
            validation = review.data_validations.dataValidation[0]
            self.assertEqual(str(validation.sqref), "E5:E34")
            self.assertEqual(validation.type, "list")
            self.assertEqual(validation.formula1, '"Correcto,Incorrecto,A revisar"')
            self.assertTrue(validation.showErrorMessage)
            rules = review.conditional_formatting[next(iter(review.conditional_formatting))]
            self.assertEqual(len(rules), 3)
            self.assertTrue(all(rule.dxf.fill.fgColor == rule.dxf.fill.bgColor for rule in rules))
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        self.assertNotIn("Objetivos y comisiones", cell.value)
            self.assertEqual(workbook["Colocacion Core"]["I5"].value, 1000000)
            self.assertEqual(workbook["Colocacion Core"]["J5"].value, 1200000)
            paths = Path(result["latest_path"]).parent / "datos"
            meta, months = SQLiteDatasetStore(next(paths.glob("*.sqlite"))).load_dataset()
            self.assertEqual(len(months), 4)
            self.assertNotIn("https://", meta.base_url)
            audit = json.loads(next(paths.glob("*.json")).read_text(encoding="utf-8"))
            self.assertIsNone(audit["manual_commission"])
            self.assertEqual(audit["manual_rules"]["high_min"], 28)
            self.assertNotIn("2026-07-10", audit["calendar"])
            self.assertEqual(sum(Decimal(row["amount"]) for row in audit["commissions"]), Decimal(3500))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
