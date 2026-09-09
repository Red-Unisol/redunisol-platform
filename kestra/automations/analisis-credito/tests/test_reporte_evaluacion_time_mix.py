from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import Mock
import random
import json

from openpyxl import Workbook

from reporte_evaluacion_comisiones.time_mix import (
    aggregate, analyze_time_mix, classify, decompose, fetch_line_context,
)
from reporte_evaluacion_comisiones.time_mix_excel import build_time_mix_sheet


def group(n, mean):
    return {"count": n, "minutes": n * mean}


def context():
    return {"applications": {1: 10, 2: 11}, "lines": {
        10: {"id": 10, "name": "Línea cancelaciones", "parent_id": 50, "parent_name": "Superior", "first_activity": "2026-07-20"},
        11: {"id": 11, "name": "Otra", "parent_id": 51, "parent_name": "Superior", "first_activity": "2020-01-01"},
    }}


class TimeMixTests(TestCase):
    def test_pure_mix_and_pure_performance(self):
        a = {"fast": group(80, 10), "slow": group(20, 30)}
        mix = decompose(a, {"fast": group(20, 10), "slow": group(80, 30)})
        self.assertAlmostEqual(mix["change"], 12)
        self.assertAlmostEqual(mix["mix"], 12)
        self.assertEqual(mix["performance"], 0)
        perf = decompose(a, {"fast": group(80, 15), "slow": group(20, 35)})
        self.assertAlmostEqual(perf["performance"], 5)
        self.assertEqual(perf["mix"], 0)

    def test_simultaneous_changes_entries_exits_and_symmetry(self):
        rng = random.Random(84)
        for _ in range(40):
            a = {k: group(rng.randint(1, 200), rng.random() * 400) for k in ["a", "b", "exit"]}
            b = {k: group(rng.randint(1, 200), rng.random() * 400) for k in ["a", "b", "entry"]}
            forward, backward = decompose(a, b), decompose(b, a)
            self.assertAlmostEqual(forward["change"], forward["mix"] + forward["performance"] + forward["entry_exit"])
            for field in ("change", "mix", "performance", "entry_exit"):
                self.assertAlmostEqual(forward[field], -backward[field])
            entry = next(r for r in forward["details"] if r["key"] == "entry")
            self.assertIsNone(entry["mean_before"])
            self.assertEqual(entry["performance"], 0)

    def test_missing_month_has_no_decomposition_instead_of_zero(self):
        for a, b in [({}, {}), ({}, {"a": group(10, 20)}), ({"a": group(10, 20)}, {})]:
            result = decompose(a, b)
            for field in ("change", "mix", "performance", "entry_exit"):
                self.assertIsNone(result[field])

    def test_new_is_first_calendar_month_and_cancellation_is_independent(self):
        c = context()
        row = {"solicitud_oid": 1, "linea": "Línea cancelaciones"}
        self.assertEqual(classify(row, c, "2026-07")[2:], ("Cancelaciones", "Nueva · primer mes"))
        self.assertEqual(classify(row, json.loads(json.dumps(c)), "2026-07"), classify(row, c, "2026-07"))
        self.assertEqual(classify(row, c, "2026-08")[2:], ("Cancelaciones", "Existente"))
        self.assertEqual(classify(row, c, "2026-06")[-1], "Antigüedad desconocida")
        c["lines"][10]["first_activity"] = None
        self.assertEqual(classify(row, c, "2026-08")[-1], "Antigüedad desconocida")

    def test_same_parent_name_does_not_merge_ids_and_missing_mapping_keeps_cases(self):
        c = context()
        rows = [{"solicitud_oid": 1, "linea": "Línea cancelaciones", "minutos": 5},
                {"solicitud_oid": 2, "linea": "Otra", "minutos": 10},
                {"solicitud_oid": 99, "linea": "Falta", "minutos": 30}]
        grouped = aggregate(rows, c, "2026-08")
        self.assertEqual(len(grouped), 3)
        self.assertEqual(sum(v["count"] for v in grouped.values()), 3)
        self.assertEqual(sum(v["minutes"] for v in grouped.values()), 45)
        rows[0]["linea"] = "Renombrada"
        self.assertEqual(classify(rows[0], c, "2026-08")[0], "Sin clasificar")
        with self.assertRaises(ValueError):
            aggregate(rows + rows[:1], c, "2026-08")

    def test_mapping_batches_strict_validation_and_global_minimum(self):
        report = SimpleNamespace(first_response=[{"solicitud_oid": 1}], transfer=[], end_to_end=[])
        client = Mock()
        client.evaluate_list.return_value = [[1, 10, "Linea", 20, "Superior"]]
        client.evaluate.return_value = "2020-01-15T00:00:00"
        c = fetch_line_context(client, [report], "2026-08")
        self.assertEqual(c["lines"][10]["first_activity"], "2020-01-15")
        expression = client.evaluate.call_args.args[0]
        self.assertIn('.Min([Fecha])', expression)
        self.assertIn('[Fecha] < #2026-09-01#', expression)
        self.assertNotIn('[Fecha] >=', expression)
        client.evaluate_list.return_value *= 2
        with self.assertRaisesRegex(ValueError, 'duplicada'):
            fetch_line_context(client, [report], "2026-08")
        client.evaluate_list.return_value = [[1, 10]]
        with self.assertRaisesRegex(ValueError, 'Esquema'):
            fetch_line_context(client, [report], "2026-08")

    def test_maturity_transition_does_not_create_artificial_entries(self):
        rows = [{"solicitud_oid": 1, "linea": "Línea cancelaciones", "minutos": 5}]
        reports = [SimpleNamespace(month_value=m, first_response=rows, transfer=rows, end_to_end=rows,
                   summary={k: {"promedio_minutos": 5} for k in ("first_response", "transfer", "end_to_end")})
                   for m in ("2026-07", "2026-08")]
        results = analyze_time_mix(reports, ["2026-08"], context())
        for result in results:
            self.assertEqual(result["entry_exit"], 0)
            self.assertEqual(result["details"][0]["key"][-1], "Existente")
        reports[0].summary["transfer"]["promedio_minutos"] = 50
        with self.assertRaisesRegex(ValueError, 'difiere'):
            analyze_time_mix(reports, ["2026-08"], context())

    def test_sheet_formula_sources_and_external_text(self):
        key = ("50", "=NO_EJECUTAR()", "Cancelaciones", "Nueva · primer mes")
        data = decompose({key: group(10, 12)}, {key: group(20, 20)})
        result = dict(month="2026-08", previous_month="2026-07", label="Primera respuesta", **data)
        wb = Workbook()
        for name in ("Comisiones", "Muestreo legajos", "Comparativo mensual"):
            wb.create_sheet(name)
        ws = build_time_mix_sheet(wb, [result], context())
        t = ws.tables["DetalleCambioTiempos"]
        from openpyxl.utils.cell import range_boundaries
        _, start, _, end = range_boundaries(t.ref)
        row = start + 1
        self.assertEqual(ws.cell(row, 3).data_type, 's')
        self.assertEqual(ws.cell(row, 16).value, 120)
        self.assertEqual(ws.cell(row, 17).value, 400)
        self.assertEqual(ws.cell(row, 10).value, f'=IF(F{row}>0,P{row}/F{row},"")')
        self.assertIn(f'G11-SUM(H11:J11)', ws['K11'].value)
        self.assertEqual(end, row)
        self.assertIsNone(ws.freeze_panes)
