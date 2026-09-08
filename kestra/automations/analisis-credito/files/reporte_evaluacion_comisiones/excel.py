from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from typing import Any, Sequence

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

from reporte_evaluacion_report.core import MonthlyReport

from .calendar import CALENDAR_DESCRIPTION, CALENDAR_VERSION, SOURCES
from .core import EXCLUDED_SELLERS, MANUAL_WEIGHT, METRICS, RULE_VERSION, Loan, previous_months

MONEY = '"$" #,##0.00'


def table_sheet(workbook: Any, name: str, title: str, subtitle: str, headers: list[str], rows: list[list], table: str) -> Any:
    ws = workbook.create_sheet(name)
    last_column = get_column_letter(len(headers))
    for row, text in [(1, title), (2, subtitle)]:
        ws.merge_cells(f"A{row}:{last_column}{row}")
        cell = ws.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor="17365D" if row == 1 else "EAF0F7")
        cell.font = Font(name="Calibri", size=15 if row == 1 else 11, bold=row == 1, color="FFFFFF" if row == 1 else "374151")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32 if row == 1 else 46
    for column, label in enumerate(headers, 1):
        cell = ws.cell(4, column, label)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(column)].width = 21
    ws.row_dimensions[4].height = 34
    for r, values in enumerate(rows, 5):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, float(value) if isinstance(value, Decimal) else value)
            # Nombres de vendedores/lineas son datos externos, nunca formulas.
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, (float, Decimal)):
                cell.number_format = "0.00"
    if rows:
        excel_table = Table(displayName=table, ref=f"A4:{last_column}{4 + len(rows)}")
        excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(excel_table)
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    return ws


def enrich_workbook(
    workbook: Any, *, reports: Sequence[MonthlyReport], months: Sequence[str],
    loans: dict[str, list[Loan]], calendar: dict, extracted_at: datetime,
) -> None:
    by_month = {report.month_value: report for report in reports}
    reference_months = sorted(by_month)
    metrics_rows = []
    for month in reference_months:
        report = by_month[month]
        metrics_rows.append([
            month, report.summary["solicitudes_count"],
            *[report.summary[group][statistic] for group, statistic, _, _ in METRICS],
        ])
    refs = table_sheet(
        workbook, "Metricas referencia", "Métricas para calcular los objetivos",
        "Valores sin redondear para el cálculo. Cada objetivo usa la media simple de los tres meses anteriores, con el mismo calendario laboral.",
        ["Mes", "Solicitudes analizadas", *[label for _, _, label, _ in METRICS]], metrics_rows, "MetricasReferencia",
    )
    ref_row = {month: row for row, month in enumerate(reference_months, 5)}
    for row in range(5, refs.max_row + 1):
        for col in range(3, 7):
            refs.cell(row, col).number_format = "0.0000"

    placement_rows = []
    placement_ranges = {}
    for month in months:
        start = 5 + len(placement_rows)
        for loan in loans[month]:
            placement_rows.append([
                month, loan.loan_id, loan.account, loan.application_id, loan.issued_on,
                loan.seller, loan.line, loan.state, loan.amount, loan.capital,
            ])
        placement_ranges[month] = (start, 4 + len(placement_rows))
    placements = table_sheet(
        workbook, "Colocacion Core", "Base de colocación del Core",
        "Préstamos emitidos en el mes con solicitud Pagada, excluyendo los vendedores indicados en Reglas. Base de comisión: Monto Deseado (MontoADesembolsar).",
        ["Mes", "ID préstamo", "Número préstamo", "Solicitud", "Fecha emisión", "Vendedor", "Línea", "Estado", "Monto para comisión", "Capital original"],
        placement_rows, "ColocacionCore",
    )
    placements.column_dimensions["F"].width = 27
    placements.column_dimensions["G"].width = 36
    for row in range(5, placements.max_row + 1):
        for col in (9, 10):
            placements.cell(row, col).number_format = MONEY

    objective_rows = []
    for month in months:
        baseline = previous_months(month)
        for _, _, label, weight in METRICS:
            objective_rows.append([month, label, baseline[0], baseline[-1], None, None, None, None, None, None, None, None, weight, None, None, None, None])
    objectives = table_sheet(
        workbook, "Objetivos y comisiones", "Objetivos y comisiones automáticas",
        "Menor tiempo es mejor. Hasta 100%: 0,5%; más de 100% y hasta 110%: 0,3%; más de 110%: 0,1%. Sin referencia positiva o datos completos: pendiente.",
        ["Mes", "Métrica", "Referencia desde", "Referencia hasta", "Mes -3 (min)", "Mes -2 (min)", "Mes -1 (min)", "Referencia (min)", "Objetivo 100% (min)", "Límite 110% (min)", "Resultado (min)", "Resultado / referencia", "Peso", "Tasa", "Colocación", "Comisión", "Estado"],
        objective_rows, "ObjetivosComisiones",
    )
    objectives.column_dimensions["B"].width = 35
    objectives.column_dimensions["Q"].width = 28
    summary_rows = [[month, len(loans[month]), None, None, None, None, None, None, float(MANUAL_WEIGHT), "Pendiente de revisión humana", None] for month in months]
    summary = table_sheet(
        workbook, "Comisiones", "Reporte de evaluación y comisiones — v2",
        "El subtotal automático corresponde al 70% del esquema. Legajos (30%) requiere revisión humana; el total definitivo queda pendiente. Importes en pesos.",
        ["Mes", "Préstamos", "Colocación", "Mediana respuesta", "Promedio respuesta", "Mediana transferencia", "Promedio transferencia", "Subtotal automático (70%)", "Peso legajos", "Estado legajos", "Total definitivo"],
        summary_rows, "ResumenComisiones",
    )
    for index, month in enumerate(months):
        sr = 5 + index
        start, end = placement_ranges[month]
        summary.cell(sr, 3, f"=SUM('Colocacion Core'!I{start}:I{end})" if end >= start else "=0")
        for metric_index, (_, _, _, weight) in enumerate(METRICS):
            row = 5 + index * 4 + metric_index
            source_col = get_column_letter(3 + metric_index)
            for col, prior in enumerate(previous_months(month), 5):
                source = f"'Metricas referencia'!{source_col}{ref_row[prior]}"
                objectives.cell(row, col, f'=IF(ISNUMBER({source}),{source},"")')
            source = f"'Metricas referencia'!{source_col}{ref_row[month]}"
            objectives.cell(row, 8, f'=IF(COUNT(E{row}:G{row})=3,AVERAGE(E{row}:G{row}),"")')
            objectives.cell(row, 9, f'=IF(ISNUMBER(H{row}),H{row},"")')
            objectives.cell(row, 10, f'=IF(ISNUMBER(H{row}),H{row}*\'Reglas\'!$B$7,"")')
            objectives.cell(row, 11, f'=IF(ISNUMBER({source}),{source},"")')
            valid = f'AND(COUNT(H{row},K{row})=2,H{row}>0)'
            objectives.cell(row, 12, f'=IF({valid},K{row}/H{row},"")')
            objectives.cell(row, 14, f'=IF({valid},IF(K{row}<=I{row},\'Reglas\'!$B$8,IF(K{row}<=J{row},\'Reglas\'!$B$9,\'Reglas\'!$B$10)),"")')
            objectives.cell(row, 15, f"='Comisiones'!C{sr}")
            objectives.cell(row, 16, f'=IF(ISNUMBER(N{row}),ROUND(O{row}*M{row}*N{row},2),"")')
            objectives.cell(row, 17, f'=IF(ISNUMBER(N{row}),"Calculado","Sin datos suficientes")')
            source = f"'Objetivos y comisiones'!P{row}"
            summary.cell(sr, 4 + metric_index, f'=IF(ISNUMBER({source}),{source},"")')
        summary.cell(sr, 8, f'=IF(COUNT(D{sr}:G{sr})=4,SUM(D{sr}:G{sr}),"")')
        summary.cell(sr, 11).comment = Comment("Pendiente: el reporte no evalúa legajos ni determina la comisión manual.", "Red Unisol")
    for row in range(5, objectives.max_row + 1):
        for col in range(5, 12):
            objectives.cell(row, col).number_format = "0.0000"
        for col in (12, 13, 14):
            objectives.cell(row, col).number_format = "0.00%"
        for col in (15, 16):
            objectives.cell(row, col).number_format = MONEY
    for row in range(5, summary.max_row + 1):
        for col in range(3, 9):
            summary.cell(row, col).number_format = MONEY
        summary.cell(row, 9).number_format = "0%"
    summary.column_dimensions["H"].width = 25
    summary.column_dimensions["J"].width = 31

    rules_rows = [
        ["Versión de reglas", RULE_VERSION],                          # B5
        ["Objetivo / referencia", 1.0],                            # B6
        ["Límite intermedio / referencia", 1.10],                   # B7
        ["Tasa hasta 100% (inclusive)", 0.005],                      # B8
        ["Tasa más de 100% hasta 110% (inclusive)", 0.003],          # B9
        ["Tasa más de 110%", 0.001],                               # B10
        ["Referencia", "Media simple de los tres meses anteriores, sin redondear antes de clasificar."],
        ["Pesos", "Respuesta: 20% mediana + 20% promedio. Transferencia: 15% mediana + 15% promedio. Legajos: 30% manual."],
        ["Sin datos / referencia cero", "Comisión pendiente, no se asigna tasa ni se muestra un total definitivo."],
        ["Legajos", "Hasta 30 solicitudes pagadas por mes, al azar con semilla reproducible. Revisión exclusivamente humana."],
        ["Calendario", CALENDAR_DESCRIPTION],
        ["Versión de calendario", CALENDAR_VERSION],
        ["Extracción iniciada (Buenos Aires)", extracted_at.isoformat()],
        ["Base monetaria", "MontoADesembolsar de préstamos emitidos en el mes con Solicitud.Estado.Descripcion = Pagada."],
        ["Vendedores excluidos", "; ".join(EXCLUDED_SELLERS)],
        ["Separación de universos", "Las métricas conservan las exclusiones de líneas del reporte original. La colocación aplica los filtros de vendedores del Core."],
        ["Histórico", "Cada ejecución conserva la base consultada. Una consulta posterior puede cambiar importes o estados de meses cerrados."],
        *[["Fuente del calendario", url] for url in SOURCES],
    ]
    rules = table_sheet(workbook, "Reglas", "Criterios y trazabilidad", "Calendario y parámetros aplicados a esta ejecución.", ["Concepto", "Valor"], rules_rows, "ReglasComisiones")
    rules.column_dimensions["A"].width = 47
    rules.column_dimensions["B"].width = 105
    for row in range(5, rules.max_row + 1):
        rules.row_dimensions[row].height = 34
    for row in range(6, 11):
        rules.cell(row, 2).number_format = "0.00%"
    table_sheet(workbook, "Feriados nacionales", "Feriados nacionales obligatorios descontados", CALENDAR_DESCRIPTION,
                ["Fecha", "Descripción"], [[day.isoformat(), name] for day, name in calendar.items()], "FeriadosNacionales")
    workbook["Feriados nacionales"].column_dimensions["B"].width = 90

    overview = workbook["Resumen ejecutivo"]
    overview["A1"] = "Reporte evaluatorio comercial — v2 con feriados nacionales"
    overview["A2"] = f"Generado: {extracted_at:%d/%m/%Y %H:%M} | {CALENDAR_DESCRIPTION}"
    overview.row_dimensions[2].height = 42
    for row in overview.iter_rows(min_col=1, max_col=2):
        if str(row[0].value).startswith(("Regla de", "Definicion de punta")):
            row[1].value = f"{row[1].value}. Se excluyen feriados nacionales obligatorios."
    sample = workbook["Muestreo legajos"]
    sample["A2"] = "Muestra reproducible de hasta 30 solicitudes pagadas por mes. Revisión manual; no se asigna puntaje ni comisión de legajos."
    sample.row_dimensions[2].height = 34
    workbook.move_sheet(summary, offset=-workbook.index(summary))
    workbook.move_sheet(objectives, offset=1 - workbook.index(objectives))
    workbook.active = 0
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
