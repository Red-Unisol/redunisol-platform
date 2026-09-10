"""Secondary trace sheet: calendar decisions are not promoted to the dashboard."""
from datetime import date

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def build_operational_calendar_sheet(workbook, audit):
    ws = workbook.create_sheet("Jornadas evaluación")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.tabColor = "A6A6A6"
    for i, width in enumerate([14, 29, 12, 12, 29, 12, 25, 16, 26, 66], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    def cell(row, col, value):
        c = ws.cell(row, col, value)
        if isinstance(value, str):
            c.data_type = "s"
        c.font = Font(name="Calibri", size=11, color="243746")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if isinstance(value, date):
            c.number_format = "yyyy-mm-dd"
        return c

    def band(row, text, title=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        c = cell(row, 1, text)
        c.fill = PatternFill("solid", fgColor="17365D" if title else "EAF0F7")
        c.font = Font(name="Calibri", size=14 if title else 11, bold=title, color="FFFFFF" if title else "243746")
        ws.row_dimensions[row].height = 30

    band(1, "Jornadas de evaluación · trazabilidad del calendario", True)
    band(2, audit["scope"] + ". Transferencia, punta a punta y duración por estado conservan el calendario nacional.")
    band(3, "Solo los cierres confirmados agregan descuentos. Ausencia de actividad = candidato a revisar; una falla de API nunca equivale a un día no trabajado.")
    band(4, "Usuarios de referencia: " + ", ".join(audit["users"]) + ". " + audit["roster_status"] + ".")
    band(5, "Señales de decisión: " + ", ".join(audit["decision_states"]) + ". No son solicitudes únicas. Se cuenta cualquier movimiento del equipo para no ignorar trabajo sin decisiones.")
    band(6, audit["date_rule"])

    def table(row, items, name):
        headers = ["Fecha", "Detección", "Eventos equipo", "Decisiones equipo", "Usuarios observados", "Eventos Core", "Cobertura", "Descuento extra", "Calendario base", "Motivo / interpretación"]
        for col, value in enumerate(headers, 1):
            c = cell(row, col, value)
            c.fill = PatternFill("solid", fgColor="305496")
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws.row_dimensions[row].height = 34
        for r, item in enumerate(items, row + 1):
            values = [date.fromisoformat(item["date"]), item["status"], item["team_events"], item["decision_events"],
                      ", ".join(item["actors"]), item["events"], item["coverage"], "Sí · 1ª respuesta" if item["extra_excluded"] else "No", item["base"], item["reason"]]
            for col, value in enumerate(values, 1):
                cell(r, col, value)
            ws.row_dimensions[r].height = 43
        if items:
            t = Table(displayName=name, ref=f"A{row}:J{row+len(items)}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        return row + len(items)

    exceptions = [r for r in audit["days"] if r["extra_excluded"] or r["status"] in
                  {"Posible no trabajado", "Sin datos completos", "Actividad por atribuir", "Confirmación en conflicto"}]
    band(8, "Excepciones y fechas que requieren revisión", True)
    end = table(9, exceptions, "ExcepcionesJornadasEvaluacion")
    if not exceptions:
        band(10, "No se detectaron excepciones en el período.")
        end = 10
    band(end + 3, "Histórico diario · no modifica automáticamente el calendario", True)
    table(end + 4, audit["days"], "HistoriaJornadasEvaluacion")
    ws.freeze_panes = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws
