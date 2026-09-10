"""Formula-driven worksheet for the mean-time decomposition."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from reporte_evaluacion_report.core import normalize_text

NAME = "Cambios en tiempos"
DATA_NAME = "Soporte tiempos"


def build_time_mix_sheet(workbook, results, context):
    from .time_mix_dashboard import build_dashboard
    data = build_time_mix_data_sheet(workbook, results, context)
    return build_dashboard(workbook, results, data)


def build_time_mix_data_sheet(workbook, results, context):
    ws = workbook.create_sheet(DATA_NAME)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 75
    widths = [14, 23, 37, 22, 24, 14, 14, 14, 14, 16, 16, 16, 16, 16, 16, 18, 18, 23]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    def cell(row, col, value, formula=False):
        c = ws.cell(row, col, value)
        if isinstance(value, str) and not formula:
            c.data_type = "s"
        c.font = Font(name="Calibri", size=11, color="243746")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if formula or isinstance(value, float):
            c.number_format = '0.00;[Red]-0.00;"—"'
        return c

    def band(row, title, dark=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
        c = cell(row, 1, title)
        c.fill = PatternFill("solid", fgColor="17365D" if dark else "EAF0F7")
        c.font = Font(name="Calibri", size=15 if dark else 11, bold=dark, color="FFFFFF" if dark else "243746")
        ws.row_dimensions[row].height = 32 if dark else 36

    def table(start, headers, end, name):
        for i, label in enumerate(headers, 1):
            c = cell(start, i, label)
            c.fill = PatternFill("solid", fgColor="305496")
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws.row_dimensions[start].height = 44
        if end > start:
            t = Table(displayName=name, ref=f"A{start}:{get_column_letter(len(headers))}{end}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)

    band(1, "¿Por qué cambiaron los tiempos?", True)
    band(2, "Promedios en minutos laborales, comparados con el mes anterior. Positivo = mayor tiempo; negativo = menor tiempo. Se conservan los casos y exclusiones del reporte.")
    band(3, "Mix = cambio de participación de los grupos. Performance = cambio de tiempo dentro del grupo. Entradas/salidas = aporte de grupos sin casos en uno de los dos meses.")
    band(4, "Nueva = solo el mes de primera actividad histórica disponible en Core, no fecha de creación. Cancelaciones = nombre de línea que contiene ‘cancel’. Son atributos cruzados, sin doble conteo.")
    band(5, "La antigüedad se clasifica al mes comparado en ambos lados, para evitar efectos artificiales al madurar. Se agrupa por ID de superior y subcategorías; faltantes quedan identificados.")
    band(6, "Base anterior fija: desempeño = participación anterior × Δtiempo; mix = Δparticipación × tiempo actual. La suma, más entradas/salidas, explica el cambio total.")
    band(7, "Describe cambios, no prueba causalidad ni una curva de aprendizaje. Superior y línea son la clasificación actual de Core; ‘Sin clasificar’ conserva los casos con datos faltantes o inconsistentes.")
    band(9, "Resumen mensual · últimas comparaciones primero", True)
    ordered = sorted(results, key=lambda r: r["month"], reverse=True)
    summary_header = 10
    category_title = summary_header + len(ordered) + 3
    category_header = category_title + 2
    category_count = sum(len({tuple(item["key"][2:]) for item in r["details"]}) for r in ordered)
    band(category_title, "Lectura por cancelaciones y líneas nuevas", True)
    band(category_title + 1, "Aportes de cada subcategoría, sumados entre superiores. Una nueva de cancelaciones ocupa una sola fila. Estos aportes desglosan el resumen; no se suman nuevamente a él.")
    detail_title = category_header + category_count + 3
    detail_header = detail_title + 2
    band(detail_title, "Aportes por línea superior y subcategoría", True)
    band(detail_title + 1, "Filtrar superior, operación o antigüedad para explorar sus aportes. Casos y minutos sumados son entradas; participaciones, promedios y efectos son fórmulas. Totales de minutos: columnas P/Q.")
    detail_row = detail_header + 1
    category_row = category_header + 1
    for summary_row, result in enumerate(ordered, summary_header + 1):
        start = detail_row
        for item in result["details"]:
            r = detail_row
            parent_id, parent_name, operation, age = item["key"]
            parent_label = f"{parent_name} [{parent_id}]" if parent_id not in ("Sin clasificar", "Sin superior") else parent_name
            values = [result["month"], result["label"], parent_label, operation, age,
                      item["count_before"], item["count_after"]]
            for col, value in enumerate(values, 1):
                cell(r, col, value)
            for col, formula in {
                8: f'=IF($C${summary_row}>0,F{r}/$C${summary_row},"")',
                9: f'=IF($D${summary_row}>0,G{r}/$D${summary_row},"")',
                10: f'=IF(F{r}>0,P{r}/F{r},"")',
                11: f'=IF(G{r}>0,Q{r}/G{r},"")',
                12: f'=IF(AND($C${summary_row}>0,$D${summary_row}>0),IF(AND(F{r}>0,G{r}>0),(I{r}-H{r})*K{r},0),"")',
                13: f'=IF(AND($C${summary_row}>0,$D${summary_row}>0),IF(AND(F{r}>0,G{r}>0),(K{r}-J{r})*H{r},0),"")',
                14: f'=IF(AND($C${summary_row}>0,$D${summary_row}>0),IF(AND(F{r}>0,G{r}>0),0,Q{r}/$D${summary_row}-P{r}/$C${summary_row}),"")',
                15: f'=IF(COUNT(L{r}:N{r})=3,SUM(L{r}:N{r}),"")',
            }.items():
                cell(r, col, formula, True)
            for col in (8, 9):
                ws.cell(r, col).number_format = "0.0%"
            cell(r, 16, item["minutes_before"])
            cell(r, 17, item["minutes_after"])
            cell(r, 18, item["status"])
            ws.row_dimensions[r].height = 36
            detail_row += 1
        end = detail_row - 1
        for operation, age in sorted({tuple(item["key"][2:]) for item in result["details"]}):
            cr = category_row
            for col, value in enumerate([result["month"], result["label"], operation, age], 1):
                cell(cr, col, value)
            def subtotal(source):
                return f'SUMIFS({source}{start}:{source}{end},D{start}:D{end},C{cr},E{start}:E{end},D{cr})'
            cell(cr, 5, '=' + subtotal('F'), True)
            cell(cr, 6, '=' + subtotal('G'), True)
            for col in (5, 6):
                ws.cell(cr, col).number_format = "#,##0"
            cell(cr, 7, f'=IF(E{cr}>0,{subtotal("P")}/E{cr},"")', True)
            cell(cr, 8, f'=IF(F{cr}>0,{subtotal("Q")}/F{cr},"")', True)
            for col, source in [(9, 'L'), (10, 'M'), (11, 'N'), (12, 'O')]:
                cell(cr, col, f'=IF(AND($C${summary_row}>0,$D${summary_row}>0),{subtotal(source)},"")', True)
            ws.row_dimensions[cr].height = 36
            category_row += 1
        r = summary_row
        cell(r, 1, result["month"])
        cell(r, 2, result["label"])
        cell(r, 3, f'=SUM(F{start}:F{end})' if end >= start else '=0', True)
        cell(r, 4, f'=SUM(G{start}:G{end})' if end >= start else '=0', True)
        for col in (3, 4):
            ws.cell(r, col).number_format = "#,##0"
        for col, numerator, denominator in [(5, "P", "C"), (6, "Q", "D")]:
            formula = f'=IF({denominator}{r}>0,SUM({numerator}{start}:{numerator}{end})/{denominator}{r},"")' if end >= start else '=""'
            cell(r, col, formula, True)
        cell(r, 7, f'=IF(COUNT(E{r}:F{r})=2,F{r}-E{r},"")', True)
        for col, source in [(8, "L"), (9, "M"), (10, "N")]:
            formula = f'=IF(COUNT(E{r}:F{r})=2,SUM({source}{start}:{source}{end}),"")' if end >= start else '=""'
            cell(r, col, formula, True)
        cell(r, 11, f'=IF(COUNT(G{r}:J{r})=4,G{r}-SUM(H{r}:J{r}),"")', True)
        cell(r, 12, f'=IF(COUNT(E{r}:F{r})<>2,"Sin comparación",IF(ABS(K{r})<0.00000001,"Conciliado","REVISAR"))', True)
        ws.row_dimensions[r].height = 32
    table(summary_header, ["Mes", "Métrica", "Casos anterior", "Casos actual", "Promedio anterior", "Promedio actual", "Cambio total (min)", "Efecto mix (min)", "Performance (min)", "Entradas/salidas (min)", "Diferencia control", "Control"], summary_header + len(ordered), "ResumenCambioTiempos")
    table(category_header, ["Mes", "Métrica", "Operación", "Antigüedad al mes", "Casos anterior", "Casos actual", "Promedio anterior", "Promedio actual", "Mix (min)", "Performance (min)", "Entradas/salidas (min)", "Aporte total (min)"], category_row - 1, "SubcategoriasCambioTiempos")
    table(detail_header, ["Mes", "Métrica", "Línea superior", "Operación", "Antigüedad al mes", "Casos anterior", "Casos actual", "% anterior", "% actual", "Promedio anterior", "Promedio actual", "Mix (min)", "Performance (min)", "Entradas/salidas (min)", "Aporte total (min)", "Suma min anterior", "Suma min actual", "Comparabilidad"], detail_row - 1, "DetalleCambioTiempos")

    catalog_title = detail_row + 2
    band(catalog_title, "Trazabilidad de líneas · clasificación actual del Core", True)
    catalog_header = catalog_title + 1
    for r, info in enumerate(sorted(context["lines"].values(), key=lambda x: x["id"]), catalog_header + 1):
        first = info["first_activity"]
        values = [info["id"], info["name"], info["parent_name"], info["parent_id"],
                  first or "Desconocida", "Cancelaciones" if "cancel" in normalize_text(info["name"]) else "Otras operaciones"]
        for col, value in enumerate(values, 1):
            cell(r, col, value)
        ws.row_dimensions[r].height = 42
    table(catalog_header, ["ID línea", "Línea", "Línea superior", "ID superior", "Primera actividad", "Operación"], catalog_header + len(context["lines"]), "CatalogoLineasTiempos")
    for col, label, target in [(1, "Ir a cancelaciones / nuevas", category_title),
                                (5, "Ir al detalle por superior", detail_title),
                                (9, "Ver catálogo de líneas", catalog_title)]:
        ws.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col + 3)
        c = cell(8, col, label)
        c.hyperlink = f"#'{DATA_NAME}'!A{target}"
        c.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
    ws.row_dimensions[8].height = 26
    for ref in [f"G11:J{summary_header + len(ordered)}", f"L{detail_header + 1}:O{detail_row - 1}"]:
        if ordered and detail_row > detail_header + 1:
            ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0.00000001"], font=Font(color="9C0006")))
            ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["-0.00000001"], font=Font(color="006100")))
    ws.freeze_panes = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = "A6A6A6"
    return ws
