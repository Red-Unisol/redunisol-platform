from __future__ import annotations

from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_monthly_comparison(workbook, months, reference_rows):
    """Una fila por mes con las seis metricas principales lado a lado."""
    del workbook["Comparativo mensual"]
    ws = workbook.create_sheet("Comparativo mensual")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.sheet_format.defaultRowHeight = 19
    for col, width in zip("ABCDEFG", [18, 20, 20, 20, 20, 20, 20]):
        ws.column_dimensions[col].width = width

    def band(row, first, last, text, color="17365D", light=False):
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        cell = ws.cell(row, first, text)
        cell.font = Font(name="Calibri", size=11, bold=not light, color="374151" if light else "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(vertical="center", horizontal="left" if light or row == 1 else "center", wrap_text=True, indent=1)
        ws.row_dimensions[row].height = 25

    band(1, 1, 7, "Comparativo mensual")
    band(2, 1, 7, "Tiempos en minutos laborales. Color por columna: verde = menor tiempo; amarillo = intermedio; rojo = mayor tiempo.", "EAF0F7", True)
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[4].height = 8
    band(5, 1, 1, "Período")
    band(5, 2, 3, "Primera respuesta")
    band(5, 4, 5, "Transferencias")
    band(5, 6, 7, "Punta a punta")
    for col, label in enumerate(["Mes", *["Mediana (min)", "Promedio (min)"] * 3], 1):
        cell = ws.cell(6, col, label)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[6].height = 24
    for row, month in enumerate(months, 7):
        ws.cell(row, 1, month)
        for metric in range(6):
            source = (f"'Comisiones'!C{10 + metric * 10}" if month == months[-1] and metric < 4
                      else f"'Metricas referencia'!{get_column_letter(3 + metric)}{reference_rows[month]}")
            ws.cell(row, metric + 2, f'=IF(ISNUMBER({source}),{source},"Pendiente")').number_format = "0.00"
        for col in range(1, 8):
            cell = ws.cell(row, col)
            cell.font = Font(name="Calibri", size=11, bold=month == months[-1], color="243746")
            cell.fill = PatternFill("solid", fgColor="F2F5F9" if row % 2 else "FFFFFF")
            cell.alignment = Alignment(horizontal="left" if col == 1 else "right", vertical="center", indent=1)
    end = 6 + len(months)
    for col in "BCDEFG":
        ws.conditional_formatting.add(f"{col}7:{col}{end}", ColorScaleRule(
            start_type="min", start_color="B7D7A8", mid_type="percentile", mid_value=50,
            mid_color="FFF2CC", end_type="max", end_color="EA9999"))
    band(end + 2, 1, 7, "Punta a punta: seguimiento, sin efecto en comisiones. Percentiles y estadísticas ampliadas se conservan en las hojas de detalle.", "EAF0F7", True)
    ws.row_dimensions[end + 2].height = 30
    ws.print_area = f"A1:G{end + 2}"
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws
