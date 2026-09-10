"""Compact controlled-contribution view; all selectors recalculate in Excel."""
from collections import Counter

from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

NAME = "Cambios en tiempos"
SIGNED = '+0.00;-0.00;"—"'
NUMBER = '0.00;-0.00;"—"'


def build_dashboard(workbook, results, data):
    ws = workbook.create_sheet(NAME)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = "24618C"
    for col, width in enumerate([5, 40, 11, 11, 13, 13, 14, 14, 14, 14], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    def put(row, col, value, *, formula=False, bold=False, fill=None, fmt=None):
        c = ws.cell(row, col, value)
        if isinstance(value, str) and not formula:
            c.data_type = "s"
        c.font = Font(name="Calibri", size=11, color="243746", bold=bold)
        c.alignment = Alignment(vertical="center", wrap_text=True,
                                horizontal="right" if fmt else "left")
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            c.number_format = fmt
        return c

    def merged(row, start, end, value, *, height=26, **kwargs):
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        ws.row_dimensions[row].height = height
        return put(row, start, value, **kwargs)

    def band(row, label):
        c = merged(row, 1, 10, label, bold=True, fill="24618C")
        c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")

    def helper(row, col, value):
        c = data.cell(row, col, value)
        if isinstance(value, str) and not value.startswith("="):
            c.data_type = "s"
        return c

    def literal_helper(row, col, value):
        c = helper(row, col, value)
        c.data_type = "s"

    source = f"'{data.title}'!"
    view = f"'{NAME}'!"
    ordered = sorted(results, key=lambda r: r["month"], reverse=True)
    months = sorted({r["month"] for r in ordered}, reverse=True)
    metrics = list(dict.fromkeys(r["label"] for r in ordered))
    parents = sorted({tuple(d["key"][:2]) for r in ordered for d in r["details"]})
    names = Counter(name for _, name in parents)

    def parent_label(key):
        pid, name = key
        return name if pid in ("Sin clasificar", "Sin superior") else f"{name} [{pid}]"

    def display_label(key):
        pid, name = key
        return f"{name} [{pid}]" if names[name] > 1 or name == "Todos los superiores" else name

    title = merged(1, 1, 10, "Aporte controlado · cambios en los tiempos", height=34, bold=True, fill="17365D")
    title.font = Font(name="Calibri", size=19, bold=True, color="FFFFFF")
    merged(2, 1, 10, "Desempeño con el mix anterior fijo. Rojo = aumenta el tiempo; verde = lo reduce.", height=25)
    merged(3, 1, 1, "Mes")
    put(3, 2, months[0] if months else "Sin datos", fill="EAF3FA", bold=True)
    merged(3, 3, 4, '=IFERROR("vs. "&TEXT(EDATE(DATE(VALUE(LEFT(B3,4)),VALUE(RIGHT(B3,2)),1),-1),"yyyy-mm"),"")', formula=True)
    put(3, 5, "Métrica")
    merged(3, 6, 10, "Primera respuesta" if "Primera respuesta" in metrics else (metrics[0] if metrics else "Sin datos"), fill="EAF3FA", bold=True)

    # Summary index and numeric cells are on the support sheet, never in the view.
    for row, item in enumerate(ordered, 4):
        literal_helper(row, 20, item["month"] + "|" + item["label"])
    helper(1, 21, f'=IFERROR(MATCH({view}$B$3&"|"&{view}$F$3,T4:T{max(4, len(ordered)+3)},0),0)')
    for col, original in enumerate("CDEFGHIJ", 22):
        helper(1, col, f'=IF(U1>0,INDEX({original}11:{original}{max(11,10+len(ordered))},U1),0)')
    comparable = f'AND({source}$U$1>0,{source}$V$1>0,{source}$W$1>0)'
    for start, end, label, summary_col in [(1, 4, "CAMBIO TOTAL", "Z"), (5, 6, "DESEMPEÑO", "AB"),
                                          (7, 8, "MIX", "AA"), (9, 10, "ENTRADAS / SALIDAS", "AC")]:
        merged(5, start, end, label, fill="EAF0F7", bold=True, height=23)
        c = merged(6, start, end, f'=IF({comparable},{source}${summary_col}$1,"")',
                   formula=True, fill="EAF0F7", height=39, fmt='+0.00" min";-0.00" min";"—"')
        c.font = Font(name="Calibri", size=23, bold=True, color="243746")
    merged(8, 1, 10,
           f'=IF({source}$U$1=0,"Sin datos para esta selección",'
           f'"Promedio: "&IF({source}$V$1>0,TEXT({source}$X$1,"0.00"),"sin casos")&" → "&'
           f'IF({source}$W$1>0,TEXT({source}$Y$1,"0.00"),"sin casos")&" min  |  Casos: "&'
           f'TEXT({source}$V$1,"#,##0")&" → "&TEXT({source}$W$1,"#,##0"))', formula=True)

    def dropdown(name, col, values, target):
        for row, value in enumerate(values, 2):
            literal_helper(row, col, value)
        letter = get_column_letter(col)
        workbook.defined_names.add(DefinedName(name, attr_text=f"{source}${letter}$2:${letter}${max(2, len(values)+1)}"))
        dv = DataValidation(type="list", formula1=name, allow_blank=False)
        dv.errorTitle, dv.error, dv.showErrorMessage = "Selección inválida", "Elegí una opción de la lista.", True
        ws.add_data_validation(dv)
        dv.add(target)

    dropdown("TiemposMeses", 35, months or ["Sin datos"], "B3")
    dropdown("TiemposMetricas", 36, metrics or ["Sin datos"], "F3")
    _, first, _, last = range_boundaries(data.tables["DetalleCambioTiempos"].ref) if "DetalleCambioTiempos" in data.tables else (1, 1, 18, 2)
    first += 1

    def source_range(col):
        return f'{source}${col}${first}:${col}${max(first,last)}'

    def subtotal(col, extra=""):
        return f'SUMIFS({source_range(col)},{source_range("A")},{view}$B$3,{source_range("B")},{view}$F$3{extra})'

    helper_start = max(len(ordered) + 7, 8)
    helper_end = helper_start + max(1, len(parents)) - 1
    for r, key in enumerate(parents, helper_start):
        literal_helper(r, 20, parent_label(key))
        literal_helper(r, 21, display_label(key))
        criterion = f',{source_range("C")},T{r}'
        for col, original in [(22, "F"), (23, "G"), (24, "P"), (25, "Q"), (28, "M"), (29, "L"), (30, "N"), (31, "O")]:
            helper(r, col, '=' + subtotal(original, criterion))
        helper(r, 26, f'=IF(V{r}>0,X{r}/V{r},"")')
        helper(r, 27, f'=IF(W{r}>0,Y{r}/W{r},"")')
        helper(r, 33, f'=V{r}+W{r}')
        helper(r, 34, f'=ABS(AB{r})')
        helper(r, 32, f'=IF(AG{r}=0,"",COUNTIFS(AH{helper_start}:AH{helper_end},">"&AH{r},AG{helper_start}:AG{helper_end},">0")+COUNTIFS(AH{helper_start}:AH{r},AH{r},AG{helper_start}:AG{r},">0"))')

    band(10, "¿Qué líneas explican el cambio? · 10 mayores aportes de desempeño, en valor absoluto")
    headers = ["#", "Línea superior", "Casos\nanterior", "Casos\nactual", "Promedio\nanterior", "Promedio\nactual", "Desempeño\n(min)", "Mix\n(min)", "Entradas /\nsalidas (min)", "Aporte total\n(min)"]

    def header(row, category=False):
        for col, label in enumerate(headers, 1):
            if category and col == 2:
                label = "Tipo de operación · antigüedad"
            c = put(row, col, label, bold=True, fill="EAF0F7")
            c.alignment = Alignment(vertical="center", horizontal="left" if col == 2 else "center", wrap_text=True)
        ws.row_dimensions[row].height = 36

    header(11)
    shown = min(10, max(1, len(parents)))
    for rank in range(1, shown + 1):
        r = 11 + rank
        index = f'MATCH({rank},{source}$AF${helper_start}:$AF${helper_end},0)'
        get = lambda col: f'INDEX({source}${col}${helper_start}:${col}${helper_end},{index})'
        put(r, 1, f'=IFERROR(IF({get("U")}<>"",{rank},""),"")', formula=True)
        for col, original in [(2, "U"), (3, "V"), (4, "W"), (5, "Z"), (6, "AA"), (7, "AB"), (8, "AC"), (9, "AD"), (10, "AE")]:
            expr = get(original)
            if col in (5, 6):
                expr = f'IF({get("V" if col == 5 else "W")}>0,{expr},"")'
            elif col >= 7:
                expr = f'IF({comparable},{expr},"")'
            put(r, col, f'=IFERROR({expr},"")', formula=True,
                fmt=None if col == 2 else "#,##0" if col in (3, 4) else SIGNED if col >= 7 else NUMBER)
        ws.row_dimensions[r].height = 28

    total_row = 12 + shown
    if len(parents) > shown:
        r = total_row
        put(r, 2, "Resto de superiores")

        def rest_sum(original):
            return f'SUMIF({source}$AF${helper_start}:$AF${helper_end},">{shown}",{source}${original}${helper_start}:${original}${helper_end})'

        for col, original in [(3, "V"), (4, "W"), (7, "AB"), (8, "AC"), (9, "AD"), (10, "AE")]:
            expr = rest_sum(original)
            if col >= 7:
                expr = f'IF({comparable},{expr},"")'
            put(r, col, '=' + expr, formula=True, fmt="#,##0" if col < 7 else SIGNED)
        for col, original, denominator in [(5, "X", "C"), (6, "Y", "D")]:
            put(r, col, f'=IF({denominator}{r}>0,{rest_sum(original)}/{denominator}{r},"")', formula=True, fmt=NUMBER)
        ws.row_dimensions[r].height = 28
        total_row += 1

    def total(row, start, end):
        put(row, 2, "TOTAL", bold=True, fill="EAF0F7")
        for col in (3, 4, 7, 8, 9, 10):
            letter = get_column_letter(col)
            expr = f'SUM({letter}{start}:{letter}{end})'
            if col >= 7:
                expr = f'IF({comparable},{expr},"")'
            put(row, col, '=' + expr, formula=True, bold=True, fill="EAF0F7", fmt="#,##0" if col < 7 else SIGNED)
        ws.row_dimensions[row].height = 28

    total(total_row, 12, total_row - 1)
    category_title = total_row + 2
    band(category_title, "Cancelaciones y líneas nuevas · detalle del aporte")
    selector_row = category_title + 1
    merged(selector_row, 1, 2, "Ver subcategorías de:")
    merged(selector_row, 3, 10, "Todos los superiores", fill="EAF3FA", bold=True)
    dropdown("TiemposSuperiores", 37, ["Todos los superiores"] + [display_label(k) for k in parents], f"C{selector_row}")
    merged(selector_row + 1, 1, 10, "Desglose del ranking, no un aporte adicional. Los minutos siempre se refieren al cambio del promedio global.", height=24)
    header(selector_row + 2, category=True)
    categories = [("Otras operaciones", "Existente"), ("Cancelaciones", "Existente"),
                  ("Otras operaciones", "Nueva · primer mes"), ("Cancelaciones", "Nueva · primer mes")]
    categories += sorted({tuple(d["key"][2:]) for item in ordered for d in item["details"]} - set(categories))
    for offset, (operation, age) in enumerate(categories):
        r = selector_row + 3 + offset
        helper_r = helper_end + 5 + offset
        literal_helper(helper_r, 20, operation)
        literal_helper(helper_r, 21, age)
        extra = f',{source_range("D")},{source}$T${helper_r},{source_range("E")},{source}$U${helper_r}'
        parent_match = f'INDEX({source}$T${helper_start}:$T${helper_end},MATCH($C${selector_row},{source}$U${helper_start}:$U${helper_end},0))'

        def category_sum(original):
            all_groups = subtotal(original, extra)
            selected = subtotal(original, extra + f',{source_range("C")},{parent_match}')
            return f'IF($C${selector_row}="Todos los superiores",{all_groups},{selected})'

        label_age = "nueva" if age == "Nueva · primer mes" else "existente" if age == "Existente" else age.lower()
        put(r, 2, f'{operation} · {label_age}')
        for col, original in [(3, "F"), (4, "G"), (7, "M"), (8, "L"), (9, "N"), (10, "O")]:
            expr = category_sum(original)
            if col >= 7:
                expr = f'IF({comparable},{expr},"")'
            put(r, col, '=' + expr, formula=True, fmt="#,##0" if col < 7 else SIGNED)
        for col, raw, denominator in [(5, "P", "C"), (6, "Q", "D")]:
            put(r, col, f'=IF({denominator}{r}>0,{category_sum(raw)}/{denominator}{r},"")', formula=True, fmt=NUMBER)
        ws.row_dimensions[r].height = 28
    category_total = selector_row + 3 + len(categories)
    total(category_total, selector_row + 3, category_total - 1)
    check_row = category_total + 2
    merged(check_row, 1, 10,
           f'=IF(NOT({comparable}),"Sin comparación: falta un mes con casos.",'
           f'IF(AND(ABS(J{total_row}-$A$6)<0.00000001,ABS(SUM(G{total_row}:I{total_row})-$A$6)<0.00000001),'
           '"Control: los aportes concilian con el cambio total.","REVISAR: los aportes no concilian."))', formula=True)
    merged(check_row + 1, 1, 10, "Nueva = solo el primer mes de actividad en Core. Cancelaciones = ‘cancel’ en el nombre. Clasificación actual; no prueba causal de aprendizaje.", height=30)
    merged(check_row + 2, 1, 10, "Sin casos en uno de los meses no hay desempeño comparable: el aporte se muestra en Entradas/salidas. Los efectos se calculan dentro de superior × operación × antigüedad.", height=30)
    c = merged(check_row + 3, 1, 10, "Ver metodología, fórmulas y datos de soporte →", height=23)
    c.hyperlink = f"#'{data.title}'!A1"
    c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    for col, explanation in [(7, "Participación del mes anterior × cambio del promedio, dentro de cada subgrupo. Se suman esos aportes por superior; no se usa el cambio del promedio agregado de la superior."),
                             (8, "Cambio de participación × promedio actual de cada subgrupo comparable."),
                             (9, "Aporte ponderado de grupos sin casos en uno de los dos meses. No se inventa un tiempo de referencia para líneas nuevas.")]:
        ws.cell(11, col).comment = Comment(explanation, "Santiago")
    for ref in [f"G12:J{total_row}", f"G{selector_row+3}:J{category_total}", "A6:J6"]:
        for operator, color, fill in [("greaterThan", "9C0006", "FCEBEC"), ("lessThan", "006100", "E9F4ED")]:
            threshold = "0.00000001" if operator == "greaterThan" else "-0.00000001"
            ws.conditional_formatting.add(ref, CellIsRule(operator=operator, formula=[threshold],
                                                        font=Font(color=color), fill=PatternFill("solid", fgColor=fill)))
    for row in list(range(12, total_row)) + list(range(selector_row+3, category_total)):
        for col in range(1, 11):
            ws.cell(row, col).border = Border(bottom=Side(style="hair", color="E4E9EF"))
    ws.freeze_panes = None
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:J{check_row+3}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    # Hide only mechanical selectors/helpers; the complete audit tables stay visible.
    data.column_dimensions.group("T", "AK", hidden=True)
    workbook.move_sheet(ws, offset=3 - workbook.index(ws))
    return ws
